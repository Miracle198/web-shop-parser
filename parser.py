"""Simple synchronous parser for the "Books to Scrape" practice store.

Downloads the catalog page, extracts book data and saves it to an Excel file.
"""

import sys
from pathlib import Path

import pandas as pd
import requests
from bs4 import BeautifulSoup
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BASE_URL = "https://books.toscrape.com/"
OUTPUT_FILE = Path(__file__).resolve().parent / "products.xlsx"
REQUEST_TIMEOUT = 10  # seconds


def fetch_page(url: str) -> str:
    """Download a page and return its HTML as a string."""
    response = requests.get(url, timeout=REQUEST_TIMEOUT)
    response.raise_for_status()
    # The server does not declare the charset, so force UTF-8
    # to avoid broken currency symbols (e.g. "Â£" instead of "£").
    response.encoding = "utf-8"
    return response.text


def parse_books(html: str) -> list[dict]:
    """Extract title, price and availability from every book card."""
    soup = BeautifulSoup(html, "html.parser")
    books = []

    for card in soup.find_all("article", class_="product_pod"):
        # The visible link text is truncated, the full title is in the attribute.
        title = card.h3.a["title"]
        price = card.find("p", class_="price_color").get_text(strip=True)
        availability = card.find("p", class_="availability").get_text(strip=True)

        books.append(
            {
                "Title": title,
                "Price": price,
                "Availability": availability,
            }
        )

    return books


def save_to_excel(books: list[dict], path: Path) -> None:
    """Save the collected data to a nicely formatted .xlsx file."""
    df = pd.DataFrame(books)

    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name="Books")
        sheet = writer.sheets["Books"]

        # Style the header row.
        header_fill = PatternFill("solid", fgColor="4F81BD")
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")

        # Auto-fit column widths based on the longest value in each column.
        for idx, column in enumerate(df.columns, start=1):
            max_len = max(df[column].astype(str).map(len).max(), len(column))
            sheet.column_dimensions[get_column_letter(idx)].width = max_len + 3

        # Keep the header visible while scrolling and enable filtering.
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions


def main() -> None:
    print(f"[INFO] Downloading catalog page: {BASE_URL}")
    try:
        html = fetch_page(BASE_URL)
    except requests.RequestException as error:
        print(f"[ERROR] Failed to download the page: {error}")
        sys.exit(1)

    books = parse_books(html)
    if not books:
        print("[WARNING] No books found. The page structure may have changed.")
        sys.exit(1)
    print(f"[INFO] Books collected: {len(books)}")

    save_to_excel(books, OUTPUT_FILE)
    print(f"[INFO] Data saved to: {OUTPUT_FILE}")


if __name__ == "__main__":
    main()
